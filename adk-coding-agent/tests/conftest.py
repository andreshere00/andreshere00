from pathlib import Path

import pytest
from docx import Document

from coding_agent.config import Settings

PROJECT_ROOT: Path = Path(__file__).resolve().parents[1]


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(
        openrouter_api_key="test-key",
        openrouter_model="test/model",
        workspace_dir=tmp_path / "ws",
        data_dir=tmp_path / "data",
        skills_dir=PROJECT_ROOT / "skills",
        sandbox_backend="process",
        sandbox_image="missing-image:latest",
        sandbox_timeout=30,
    )


def make_docx(path: Path, paragraphs: list[str], heading: str | None = None) -> Path:
    """Creates a .docx with an optional heading, the paragraphs and a one-row table."""
    document = Document()
    if heading:
        document.add_heading(heading, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Price"
    table.rows[0].cells[1].text = "100"
    document.save(str(path))
    return path

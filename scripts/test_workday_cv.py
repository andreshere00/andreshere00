#!/usr/bin/env python3
"""Checks that the Workday resume contains the fields ATS parsers expect."""

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "cv" / "Andres_Herencia_CV_Workday.docx"
TXT_PATH = ROOT / "cv" / "Andres_Herencia_CV_Workday.txt"

REQUIRED_FRAGMENTS = [
    "Andrés Herencia",
    "andresherencia2000@gmail.com",
    "+34 608 325 729",
    "Madrid, Spain",
    "WORK EXPERIENCE",
    "Forward Deployed AI Engineer",
    "InnoIT",
    "Minsait",
    "BBVA",
    "The Cliff",
    "SUNRISE",
    "Cabify",
    "February 2026 - Present",
    "EDUCATION",
    "Master's Degree",
    "CERTIFICATIONS",
    "AWS Certified Solutions Architect",
    "AWS Certified Machine Learning Engineer",
    "SKILLS",
    "Python",
    "LANGUAGES",
    "Spanish",
]


def extract_docx_text(path: Path) -> str:
    document = Document(path)
    assert not document.tables, "Workday resumes must not use tables"
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _contains(haystack: str, needle: str) -> bool:
    return needle.casefold() in haystack.casefold()


def test_docx_contains_required_ats_fields() -> None:
    text = extract_docx_text(DOCX_PATH)
    missing = [fragment for fragment in REQUIRED_FRAGMENTS if not _contains(text, fragment)]
    assert not missing, f"DOCX missing fragments: {missing}"


def test_txt_contains_required_ats_fields() -> None:
    text = TXT_PATH.read_text(encoding="utf-8")
    missing = [fragment for fragment in REQUIRED_FRAGMENTS if not _contains(text, fragment)]
    assert not missing, f"TXT missing fragments: {missing}"


if __name__ == "__main__":
    test_docx_contains_required_ats_fields()
    test_txt_contains_required_ats_fields()
    print("ATS field checks passed")
